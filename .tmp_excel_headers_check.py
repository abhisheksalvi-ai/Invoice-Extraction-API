from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

pdf_path = Path(r"c:\Users\abhis\Downloads\Rockwell - Inv Consolidation.pdf")
client = TestClient(app)

with pdf_path.open("rb") as f:
    response = client.post(
        "/download-invoices",
        files={"file": (pdf_path.name, f.read(), "application/pdf")},
    )

print(f"STATUS={response.status_code}")
if response.status_code != 200:
    print(response.text)
    raise SystemExit(1)

wb = load_workbook(filename=BytesIO(response.content), data_only=True)
ws = wb["Invoice_2"]

header_row = None
for row in ws.iter_rows(min_row=1, max_col=9, values_only=True):
    if row and row[0] == "Item No.":
        header_row = row
        break

print(f"HEADERS={header_row}")
if header_row is None:
    raise SystemExit(2)

# Print first 5 line items to verify EAN and header mapping.
seen = 0
capture = False
for row in ws.iter_rows(min_row=1, max_col=9, values_only=True):
    if row and row[0] == "Item No.":
        capture = True
        continue
    if capture:
        if row[0] in (None, "Validation"):
            break
        print(f"ROW={row}")
        seen += 1
        if seen >= 5:
            break
