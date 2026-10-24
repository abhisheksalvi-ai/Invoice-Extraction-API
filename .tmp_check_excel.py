import asyncio
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
from app.services.document_processor import DocumentProcessor
from app.services.excel_service import ExcelService

pdf_path = Path(r"c:\Users\abhis\Downloads\CUBIC - 6263375.pdf")

async def main():
    processor = DocumentProcessor()
    _, invoices = await processor.process_document(pdf_path.read_bytes())
    xlsx = ExcelService.generate_workbook_bytes(invoices)
    wb = load_workbook(BytesIO(xlsx), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=1, max_row=25, values_only=True):
        print(row)

asyncio.run(main())
