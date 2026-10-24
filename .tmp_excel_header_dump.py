from pathlib import Path
from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app

pdf = Path('sample_pdfs') / 'PROSOFT - 04-2602846.pdf'
client = TestClient(app)
resp = client.post('/download-invoices', files={'file': (pdf.name, pdf.read_bytes(), 'application/pdf')})
print('status', resp.status_code)
wb = load_workbook(BytesIO(resp.content))
ws = wb[wb.sheetnames[0]]

# Print header fields table rows (key, value) section.
for r in range(6, 32):
    k = ws.cell(r, 1).value
    v = ws.cell(r, 2).value
    if k is not None:
        print(f"{r}: {k} => {v}")
