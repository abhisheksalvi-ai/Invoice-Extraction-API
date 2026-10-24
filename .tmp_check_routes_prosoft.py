from pathlib import Path
from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app

pdf = Path('sample_pdfs') / 'PROSOFT - 04-2602846.pdf'
client = TestClient(app)

resp = client.post('/extract-invoices', files={'file': (pdf.name, pdf.read_bytes(), 'application/pdf')})
print('extract status', resp.status_code)
if resp.status_code == 200:
    body = resp.json()
    print('invoiceCount', body.get('invoiceCount'))
    if body.get('invoices'):
        inv = body['invoices'][0]
        print('invoiceNumber', inv.get('invoiceNumber'))
        print('lineItems', len(inv.get('lineItems', [])))
        print('firstItems', inv.get('lineItems', [])[:3])
else:
    print(resp.text)

resp2 = client.post('/download-invoices', files={'file': (pdf.name, pdf.read_bytes(), 'application/pdf')})
print('download status', resp2.status_code)
if resp2.status_code == 200:
    wb = load_workbook(BytesIO(resp2.content))
    ws = wb[wb.sheetnames[0]]
    non_empty = []
    for r in range(24, 40):
        vals = [ws.cell(r, c).value for c in range(1, 10)]
        if any(v is not None and str(v).strip() != '' for v in vals):
            non_empty.append((r, vals))
    print('rows', non_empty)
else:
    print(resp2.text)
