import asyncio
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
from app.services.document_processor import DocumentProcessor
from app.services.excel_service import ExcelService

async def main():
    path = Path('sample_pdfs') / 'PROSOFT - 04-2602846.pdf'
    processor = DocumentProcessor()
    pages, invoices = await processor.process_document(path.read_bytes())
    print('pages', pages, 'invoices', len(invoices), 'line_items', sum(len(i.lineItems) for i in invoices))
    for inv in invoices:
        print('invoice', inv.invoiceNumber, 'items', len(inv.lineItems), 'total', inv.header.totalAmount)
        for li in inv.lineItems:
            print('  ', li.itemNumber, li.partNumber, li.quantity, li.unitPrice, li.amount)

    xlsx = ExcelService.generate_workbook_bytes(invoices)
    wb = load_workbook(BytesIO(xlsx))
    ws = wb[wb.sheetnames[0]]
    # print first non-empty rows around line items section
    for r in range(20, 45):
        row_vals = [ws.cell(r, c).value for c in range(1, 11)]
        if any(v is not None and str(v).strip() != '' for v in row_vals):
            print('ROW', r, row_vals)

asyncio.run(main())
